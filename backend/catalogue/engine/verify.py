"""Independent check of a finished IPRS reconciliation.

The generated master is read back from the file itself and compared with the uploaded master cell by cell.
Every changed cell must be explained (an IPRS block written from the IPRS report, a value the reviewer chose,
or a new song row); anything else is an unexpected change. Control totals are checked for every step, and the
whole sheet is returned as a before/after preview.
"""
import io
from collections import Counter, defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

from .fill import AMEND, NOT_OUR_WORK, NOT_REGISTERED, UPLOADED, is_not_our_work
from .master import read_layout
from .normalize import clean_text, norm_isrc, norm_iswc, norm_title, same_party, title_similarity

PASS, WARN, FAIL, INFO = "pass", "warn", "fail", "info"
STATUSES = {UPLOADED, AMEND, NOT_REGISTERED, NOT_OUR_WORK}
EDIT_FIELD = {"title": "Title", "duration": "Duration", "iswc": "ISWC", "isrc": "ISRC"}
SLOT_GROUP = {"composers": "Composers", "authors": "Authors", "publishers": "Publishers"}


def _text(v):
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return clean_text(v)


def _rows(ws):
    """Data rows: rows 2.. that have an SN or a title."""
    return [r for r in range(2, ws.max_row + 1) if any(ws.cell(r, c).value not in (None, "") for c in (1, 2))]


def _columns(layout, ncols, header):
    """col index (0-based) -> {group, role, block}."""
    cols = {}
    for name, idx in layout["fields"].items():
        cols[idx] = {"group": "Work details", "role": name}
    for kind, slots in layout["slots"].items():
        for n, slot in enumerate(slots, start=1):
            for role, idx in slot.items():
                cols[idx] = {"group": kind.title(), "role": role, "slot": n}
    for b in layout["blocks"]:
        label = f"{b['society']} {b['slot']}"
        for role in ("isrc", "iswc", "code"):
            cols[b[role]] = {"group": label, "role": role, "block": b}
        for role, idx in b["fields"].items():
            cols[idx] = {"group": label, "role": role, "block": b}
    for idx in range(ncols):
        cols.setdefault(idx, {"group": "Work details" if idx < 58 else "Other", "role": _text(header[idx]) or f"column {idx + 1}"})
    return cols


def verify(master_path, out_bytes, iprs, mapping, final, fill, tunecodes, edits, client, amend_bytes=None, reg_bytes=None):
    before_wb = openpyxl.load_workbook(master_path)
    after_wb = openpyxl.load_workbook(io.BytesIO(out_bytes))
    before, after = before_wb.worksheets[0], after_wb.worksheets[0]
    header = [c.value for c in before[1]]
    ncols = len(header)
    layout = read_layout(header)
    cols = _columns(layout, ncols, header)
    iprs_blocks = [b for b in layout["blocks"] if b["society"] == "IPRS"]
    works = {w["internal_no"]: w for w in iprs["works"]}
    links = {l["internal_no"]: l for l in mapping["links"]}
    before_rows, after_rows = _rows(before), _rows(after)
    added_rows = [r for r in after_rows if r > (before_rows[-1] if before_rows else 1)]
    edit_by = {(e["master_row"], e["field"]): e for e in edits if e["choice"] == "iprs"}
    checks = []

    def check(step, cid, label, status, detail, items=()):
        checks.append({"step": step, "id": cid, "label": label, "status": status, "detail": detail, "items": list(items)[:500]})

    def title_of(ws, r):
        return _text(ws.cell(r, layout["fields"]["title"] + 1).value) if "title" in layout["fields"] else ""

    # ---- what the generated file actually holds in the IPRS blocks
    out_blocks = defaultdict(list)   # row -> [{slot, code, iswc, isrc, client, general, purpose}]
    code_rows = defaultdict(list)    # code -> [rows]
    for r in after_rows:
        for b in iprs_blocks:
            f = b["fields"]
            got = {"slot": b["slot"], "code": _text(after.cell(r, b["code"] + 1).value),
                   "iswc": _text(after.cell(r, b["iswc"] + 1).value), "isrc": _text(after.cell(r, b["isrc"] + 1).value),
                   "client": _text(after.cell(r, f["client_status"] + 1).value) if "client_status" in f else "",
                   "general": _text(after.cell(r, f["general_status"] + 1).value) if "general_status" in f else "",
                   "purpose": _text(after.cell(r, f["amendment_purpose"] + 1).value) if "amendment_purpose" in f else ""}
            if any(v for k, v in got.items() if k != "slot"):
                out_blocks[r].append(got)
            if got["code"]:
                code_rows[got["code"]].append(r)

    # ================= Step 1 · Upload file & check headers
    failed = [c for c in iprs["checks"] if not c["passed"]]
    check(1, "headers", "IPRS file has the fixed IPRS columns", FAIL if failed else PASS,
          f"{len(iprs['checks']) - len(failed)} of {len(iprs['checks'])} format checks passed",
          [{"title": c["label"], "detail": c["detail"]} for c in failed])
    rows_read = sum(len(w["contributors"]) for w in iprs["works"])
    footer = iprs["stats"].get("footer_count")
    ok = rows_read == iprs["stats"]["contributor_rows"] and (footer in (None, len(iprs["works"])))
    check(1, "rows", "Every row of the IPRS report was read", PASS if ok else FAIL,
          f"{iprs['stats']['contributor_rows']} rows → {len(iprs['works'])} works"
          + (f" (the report's footer says {footer})" if footer is not None else ""))

    # ================= Step 2 · Map with master
    states = Counter(l["state"] for l in mapping["links"])
    total = states["linked"] + states["review"] + states["unmatched"]
    check(2, "outcome", "Every IPRS work has exactly one outcome", PASS if total == len(iprs["works"]) else FAIL,
          f"{states['linked']} linked + {states['review']} waiting + {states['unmatched']} not in master = {total} of {len(iprs['works'])} works")
    waiting = [l for l in mapping["links"] if l["state"] == "review"]
    check(2, "waiting", "No match is waiting for your decision", WARN if waiting else PASS,
          f"{len(waiting)} waiting – they are not written into the master until you decide" if waiting else "All decided",
          [{"code": l["internal_no"], "title": l["title"], "row": l["candidates"][0]["row"] if l["candidates"] else None,
            "detail": f"suggested: {l['candidates'][0]['title']}" if l["candidates"] else "no suggestion"} for l in waiting])
    by_title = [l for l in mapping["links"] if l["state"] == "linked" and l["method"] in ("Title + writers", "Title + contributors", "Confirmed by reviewer")]
    check(2, "title_links", "Links made without an identifier (check by eye)", INFO,
          f"{len(by_title)} linked by title and writers or by your confirmation – no tune code / ISWC / ISRC in common",
          [{"code": l["internal_no"], "title": l["title"], "row": l["master_row"],
            "detail": f"→ {l['master_title']} · {l['method']} · title {title_similarity(l['title'], l['master_title'])}% similar"} for l in by_title])
    close = [l for l in mapping["links"] if l["state"] == "linked" and title_similarity(l["title"], l.get("master_title", "")) < 100]
    check(2, "titles", "Linked songs whose titles are not identical", INFO if close else PASS,
          f"{len(close)} – usually a version (LIVE, LYRIC VIDEO, FEAT …) or a spelling difference",
          [{"code": l["internal_no"], "title": l["title"], "row": l["master_row"], "detail": f"master: {l['master_title']}"} for l in close])
    per_row = defaultdict(list)
    for l in mapping["links"]:
        if l["state"] == "linked":
            per_row[l["master_row"]].append(l)
    multi = {r: ls for r, ls in per_row.items() if len(ls) > 1}
    check(2, "multi", "Songs registered more than once at IPRS (kept on one row)", INFO,
          f"{len(multi)} songs with 2+ IPRS works – they fill IPRS blocks 1, 2, 3 … of the same row",
          [{"row": r, "title": ls[0]["master_title"], "detail": ", ".join(f"{l['internal_no']} {l['title']}" for l in ls)} for r, ls in sorted(multi.items())])

    # ================= Step 3 · Fill IPRS columns (read back from the generated file)
    overflow = {c for f in fill["rows"] for c in f["overflow"]}
    expected = {l["internal_no"]: l for l in mapping["links"] if l["state"] in ("linked", "unmatched")}
    missing = [c for c in expected if c not in code_rows and c not in overflow]
    check(3, "present", "Every IPRS work is in the final master", FAIL if missing else PASS,
          f"{len(expected) - len(missing) - len(overflow & set(expected))} of {len(expected)} linked / new works found in an IPRS TUNECODE column"
          + (f" · {len(overflow)} did not fit (no free IPRS block)" if overflow else ""),
          [{"code": c, "title": works[c]["title"], "detail": "not found in any IPRS TUNECODE column"} for c in missing])
    dupes = {c: rs for c, rs in code_rows.items() if len(rs) > 1}
    check(3, "once", "No IPRS tune code is written twice", FAIL if dupes else PASS,
          f"{len(code_rows)} tune codes, each in one place" if not dupes else f"{len(dupes)} tune codes appear more than once",
          [{"code": c, "title": works.get(c, {}).get("title", ""), "row": rs[0], "detail": f"rows {', '.join(map(str, rs))}"} for c, rs in dupes.items()])
    wrong_row = []
    for c, l in expected.items():
        rows = code_rows.get(c, [])
        if not rows:
            continue
        want = l["master_row"] if l["state"] == "linked" else None
        if (want and rows[0] != want) or (want is None and rows[0] not in added_rows):
            wrong_row.append({"code": c, "title": l["title"], "row": rows[0],
                              "detail": f"expected {'row ' + str(want) if want else 'a new row'}, found in row {rows[0]}"})
    check(3, "right_row", "Each tune code sits in the row of the song it was linked to", FAIL if wrong_row else PASS,
          "All in the linked song's row (new songs in their own new row)" if not wrong_row else f"{len(wrong_row)} in the wrong row", wrong_row)
    id_bad, foreign = [], []
    for r, blocks in out_blocks.items():
        for b in blocks:
            w = works.get(b["code"])
            if not b["code"]:
                continue
            if not w:
                foreign.append({"code": b["code"], "row": r, "title": title_of(after, r), "detail": f"IPRS block {b['slot']} – already in the master, not in this IPRS report"})
                continue
            for label, got, want in (("ISWC", norm_iswc(b["iswc"]), w["iswc_valid"]), ("ISRC", norm_isrc(b["isrc"]), w["isrc_valid"])):
                if want and got != want:
                    id_bad.append({"code": b["code"], "row": r, "title": title_of(after, r), "detail": f"block {b['slot']} {label}: file {b[label.lower()] or '—'}, IPRS report {want}"})
    check(3, "ids", "ISWC and ISRC in each block are the IPRS report's values", FAIL if id_bad else PASS,
          "Every block's ISWC / ISRC equals its IPRS work" if not id_bad else f"{len(id_bad)} differ", id_bad)
    if foreign:
        check(3, "foreign", "Tune codes in the master that this IPRS report does not list", WARN,
              f"{len(foreign)} – kept as they were; the client may not be credited on them any more", foreign)
    status_bad, client_bad = [], []
    for r, blocks in out_blocks.items():
        for b in blocks:
            for label, value in (("Client status", b["client"]), ("General Status", b["general"])):
                if value and value not in STATUSES and not is_not_our_work(value):
                    status_bad.append({"row": r, "title": title_of(after, r), "detail": f"block {b['slot']} {label} = '{value}'"})
            amend = AMEND in (b["client"], b["general"])
            if amend and not b["purpose"]:
                status_bad.append({"row": r, "title": title_of(after, r), "detail": f"block {b['slot']} says Amend but Purpose of amendment is empty"})
            if b["purpose"] and not amend and b["client"] in (UPLOADED,) and b["general"] in (UPLOADED,):
                status_bad.append({"row": r, "title": title_of(after, r), "detail": f"block {b['slot']} is Uploaded but has a Purpose of amendment"})
            w = works.get(b["code"])
            if w and b["client"] == UPLOADED and not any(same_party(c["name"], client["name"], c["ipi"], client.get("ipi_name_no", "")) for c in w["contributors"]):
                client_bad.append({"code": b["code"], "row": r, "title": w["title"], "detail": "Client status Uploaded, but the client is not on the IPRS work"})
    check(3, "status", "Client status / General Status / Purpose of amendment agree", FAIL if status_bad else PASS,
          "Amend always has a reason; Uploaded never does" if not status_bad else f"{len(status_bad)} problems", status_bad)
    check(3, "client", "'Uploaded' only where IPRS credits the client", FAIL if client_bad else PASS,
          f"Checked against the client {client['name'].title()} (IPI {client.get('ipi_name_no') or '—'})", client_bad)
    no_code = [r for r in after_rows if not any(b["code"] for b in out_blocks.get(r, []))]
    pending_titles = {norm_title(f["title"]) for f in fill["rows"] if f["state"] == "pending"}
    unmarked = [r for r in no_code if not any(b["client"] or b["general"] for b in out_blocks.get(r, []))]
    check(3, "unregistered", "Songs with no IPRS registration are marked", WARN if unmarked else PASS,
          f"{len(no_code)} songs without an IPRS tune code – {len(no_code) - len(unmarked)} marked Not registered / Not our work",
          [{"row": r, "title": title_of(after, r), "detail": "waiting for your decision in Map with master" if norm_title(title_of(after, r)) in pending_titles else "no status written"} for r in unmarked])

    titles = Counter(norm_title(title_of(after, r)) for r in after_rows)
    dup_titles = [{"row": r, "title": title_of(after, r), "detail": f"title appears {titles[norm_title(title_of(after, r))]} times"}
                  for r in after_rows if titles[norm_title(title_of(after, r))] > 1]
    check(3, "no_dupes", "No song appears twice in the final master", WARN if dup_titles else PASS,
          f"{len(after_rows)} songs, {len(titles)} different titles" + (" – same title on several rows, check they are different songs" if dup_titles else ""), dup_titles)
    renamed = {e["master_row"] for e in edits if e["choice"] == "iprs" and e["field"] == "Title"}
    lost = [{"row": r, "title": title_of(before, r), "detail": f"now '{title_of(after, r)}'"} for r in before_rows
            if r not in renamed and norm_title(title_of(before, r)) != norm_title(title_of(after, r))]
    expect_rows = len(before_rows) + len(final["added"])
    check(3, "rows_kept", "Every song of our master is kept, in the same row", FAIL if lost or len(after_rows) != expect_rows else PASS,
          f"{len(before_rows)} master songs + {len(final['added'])} new from IPRS = {expect_rows}; the file has {len(after_rows)}", lost)

    # ---- cell-level lineage: every changed cell must be explained
    # Read each sheet once (openpyxl's max_row and per-cell lookups are slow on 1 400 columns).
    def grid(ws):
        out = {}
        for r, values in enumerate(ws.iter_rows(min_row=1, max_col=ncols, values_only=True), start=1):
            out[r] = [_text(v) for v in values] + [""] * (ncols - len(values))
        return out
    before_grid, after_grid = grid(before), grid(after)
    BLANK = [""] * ncols
    added_set = set(added_rows)
    added_code = {r: next((b["code"] for b in out_blocks.get(r, []) if b["code"]), "") for r in added_rows}
    changed, unexplained = [], []
    preview_cells = defaultdict(list)
    used_cols = set()
    for r in after_rows:
        for c in range(ncols):
            a, b = after_grid[r][c], (before_grid.get(r) or BLANK)[c] if r not in added_set else ""
            if a or b:
                used_cols.add(c)
            if a == b:
                if a:
                    preview_cells[r].append([c, b, a, "same", ""])
                continue
            info = cols[c]
            why = None
            if r in added_rows:
                w = works.get(added_code[r])
                why = f"New song from IPRS work {added_code[r]}" + (f" ({w['title']})" if w else "")
            elif info.get("block") is not None and info["block"]["society"] == "IPRS":
                code = _text(after.cell(r, info["block"]["code"] + 1).value)
                why = f"IPRS block {info['block']['slot']} · from IPRS work {code}" if code else f"IPRS block {info['block']['slot']} · status for a song IPRS does not have"
            elif info["group"] == "Work details" and EDIT_FIELD.get(info["role"]) and (r, EDIT_FIELD[info["role"]]) in edit_by:
                why = f"You chose the IPRS {EDIT_FIELD[info['role']]} in Map with master"
            elif info["group"].lower() in SLOT_GROUP and (r, SLOT_GROUP[info["group"].lower()]) in edit_by:
                why = f"You chose the IPRS {SLOT_GROUP[info['group'].lower()]} in Map with master"
            elif info.get("block") is not None and info["role"] in ("iswc", "isrc") and (r, info["role"].upper()) in edit_by:
                why = f"You chose the IPRS {info['role'].upper()} in Map with master"
            state = "changed" if why else "unexpected"
            if not why:
                unexplained.append({"row": r, "title": title_of(after, r),
                                    "detail": f"{get_column_letter(c + 1)} {_text(header[c])}: '{b}' → '{a}'"})
            changed.append(r)
            preview_cells[r].append([c, b, a, "added" if r in added_rows else state, why or "Not explained by any step"])
    check(3, "lineage", "Every changed cell is explained (nothing else was touched)", FAIL if unexplained else PASS,
          f"{len(changed)} cells changed or added – each traced to the IPRS report or to your choices" if not unexplained
          else f"{len(unexplained)} of {len(changed)} changed cells have no explanation", unexplained)

    # ================= Step 4 · Reports
    status_count = Counter(r["societies"]["IPRS"]["status"] for r in tunecodes["rows"] if "IPRS" in r["societies"])
    for cid, label, data, status in (("amend_report", "Amend report has every song to amend", amend_bytes, "Amend"),
                                     ("reg_report", "Registration report has every song to register", reg_bytes, "Registration")):
        if data is None:
            continue
        ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
        n = len(_rows(ws))
        check(4, cid, label, PASS if n == status_count[status] else FAIL,
              f"{n} rows in the file, {status_count[status]} songs with IPRS status {status}")
    check(4, "final_count", "Final master = our master + songs added from IPRS",
          PASS if len(after_rows) == final["summary"]["final_total"] else FAIL,
          f"{final['summary']['master_songs']} + {final['summary']['added']} = {final['summary']['final_total']}; file has {len(after_rows)} songs")

    # ---- preview: the used columns of the final master, before and after
    groups = []
    for c in sorted(used_cols):
        info = cols[c]
        groups.append({"col": c, "letter": get_column_letter(c + 1), "header": _text(header[c]) or info["role"], "group": info["group"]})
    preview_rows = [{"row": r, "title": title_of(after, r), "kind": "added" if r in added_set else "master",
                     "changed": sum(1 for x in preview_cells[r] if x[3] != "same"),
                     "unexpected": sum(1 for x in preview_cells[r] if x[3] == "unexpected"),
                     "amend": any(AMEND in (b["client"], b["general"]) for b in out_blocks.get(r, [])),
                     "cells": preview_cells[r]} for r in after_rows]
    summary = Counter(c["status"] for c in checks)
    return {"checks": checks, "summary": {"pass": summary[PASS], "warn": summary[WARN], "fail": summary[FAIL], "info": summary[INFO],
                                          "cells_changed": len(changed), "cells_unexplained": len(unexplained),
                                          "rows": len(after_rows), "added": len(added_rows)},
            "preview": {"columns": groups, "rows": preview_rows}}
