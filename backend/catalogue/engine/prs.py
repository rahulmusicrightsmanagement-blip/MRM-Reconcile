"""Read a PRS 'my works' export (CSV or Excel): one row per tunecode with
WRITER n / PUBLISHER n columns."""
import csv
import io
import re

import openpyxl

from .normalize import clean_id, clean_text, norm_iswc


class SocietyFormatError(Exception):
    pass


def _read_table(path):
    if str(path).lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        rows = [[clean_id(v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raw = open(path, "rb").read()
        text = raw.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(clean_text(v) for v in r)]
    if not rows:
        raise SocietyFormatError("The file is empty")
    return [clean_text(h).upper() for h in rows[0]], rows[1:]


def parse_prs(path):
    header, rows = _read_table(path)
    if "TITLE" not in header or "TUNECODE" not in header:
        raise SocietyFormatError("PRS export needs TITLE and TUNECODE columns")
    idx = {h: i for i, h in enumerate(header)}
    writer_cols = [i for i, h in enumerate(header) if re.fullmatch(r"WRITER\s*\d+", h)]
    publisher_cols = [i for i, h in enumerate(header) if re.fullmatch(r"PUBLISHER\s*\d+", h)]

    def get(row, name):
        i = idx.get(name)
        return clean_text(row[i]) if i is not None and i < len(row) else ""

    works = []
    for n, row in enumerate(rows, start=2):
        code = get(row, "TUNECODE")
        if not code:
            continue
        works.append({
            "row": n,
            "title": get(row, "TITLE"),
            "code": code,
            "iswc": get(row, "ISWC"),
            "iswc_valid": norm_iswc(get(row, "ISWC")),
            "writers": [clean_text(row[i]) for i in writer_cols if i < len(row) and clean_text(row[i])],
            "publishers": [clean_text(row[i]) for i in publisher_cols if i < len(row) and clean_text(row[i])],
            "work_status": get(row, "WORK STATUS"),
            "distributed": get(row, "PREVIOUSLY DISTRIBUTED PRS/MCPS"),
            "other_info": get(row, "OTHER WORK INFO"),
        })
    checks = [
        {"label": "TITLE and TUNECODE columns", "passed": True, "detail": "found"},
        {"label": "Writer columns", "passed": bool(writer_cols), "detail": f"{len(writer_cols)} found (WRITER 1 … {len(writer_cols)})"},
        {"label": "Publisher columns", "passed": True, "detail": f"{len(publisher_cols)} found" if publisher_cols else "none – publishers will not be compared"},
        {"label": "ISWC column", "passed": "ISWC" in idx, "detail": "found" if "ISWC" in idx else "missing – matching uses tune codes and titles only"},
        {"label": "Works read", "passed": bool(works), "detail": f"{len(works)} rows with a tune code"},
    ]
    return {"works": works, "writer_columns": len(writer_cols), "checks": checks, "stats": {"works": len(works)}}
