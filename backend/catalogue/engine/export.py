"""Excel outputs for each step. Every builder writes into a workbook so the
full report can reuse them."""
import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .iprs import SHARE_KEYS
from .master import read_layout
from .normalize import clean_text, norm_ipi, same_party

HEADER_FILL = PatternFill("solid", fgColor="1F4E46")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {
    "Match": PatternFill("solid", fgColor="D9F2E3"), "matched": PatternFill("solid", fgColor="D9F2E3"),
    "Amend": PatternFill("solid", fgColor="FFF1CC"), "differences": PatternFill("solid", fgColor="FFF1CC"),
    "Registration": PatternFill("solid", fgColor="FADBD8"), "unmatched": PatternFill("solid", fgColor="FADBD8"),
    "review": PatternFill("solid", fgColor="DDE7F7"),
    "error": PatternFill("solid", fgColor="FADBD8"), "warning": PatternFill("solid", fgColor="FFF1CC"),
}
ADDED_FILL = PatternFill("solid", fgColor="D9F2E3")
FILLED_FILL = PatternFill("solid", fgColor="FFF1CC")


def _sheet(wb, title, header, rows, status_col=None, widths=None):
    ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "-", title)[:31])  # Excel forbids these characters in sheet names
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(["" if v is None else v for v in row])
        if status_col is not None:
            fill = STATUS_FILLS.get(str(row[status_col]))
            if fill:
                ws.cell(ws.max_row, status_col + 1).fill = fill
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    for i, h in enumerate(header, start=1):
        width = (widths or {}).get(h)
        if width is None:
            sample = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)]
            width = min(max(sample) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws


def _new_book():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Step 1 -----------------------------------------------------------------
def add_step1(wb, iprs):
    _sheet(wb, "1 Validation", ["Check", "Result", "Detail"],
           [[c["label"], "Passed" if c["passed"] else "Failed", c["detail"]] for c in iprs["checks"]])
    _sheet(wb, "1 Data issues", ["Internal No", "Title", "Severity", "Issue", "Detail"],
           [[i["internal_no"], i["title"], i["severity"], i["code"], i["message"]] for i in iprs["issues"]], status_col=2)
    _sheet(wb, "1 Cleanup log", ["Excel row", "Column", "Original value", "Cleaned value", "Action"],
           [[c["row"], c["column"], c["original"], c["cleaned"], c["action"]] for c in iprs["cleanup"]])


# ---- Step 2 -----------------------------------------------------------------
def oneline_table(iprs):
    works = iprs["works"]
    counts = {g: max([len(w[g]) for w in works] + [1]) for g in ("authors", "composers", "publishers")}
    header = ["Internal No", "Title", "Alternate Title", "ISWC", "ISRC", "Duration", "Language", "Category", "Status",
              "Performer", "Production Title", "Authors", "Composers", "Publishers"]
    for group, label in (("authors", "Author"), ("composers", "Composer"), ("publishers", "Publisher")):
        for n in range(1, counts[group] + 1):
            header += [f"{label} {n}", f"{label} {n} Role", f"{label} {n} IPI", f"{label} {n} Society", f"{label} {n} PER %"]
    header += ["Issues", "Source rows"]
    rows = []
    for w in works:
        row = [w["internal_no"], w["title"], w["alt_title"], w["iswc"], w["isrc"], w["duration"], w["language"],
               w["category"], w["status"], w["performer"], w["production"],
               len(w["authors"]), len(w["composers"]), len(w["publishers"])]
        for group in ("authors", "composers", "publishers"):
            people = w[group]
            for n in range(counts[group]):
                if n < len(people):
                    c = people[n]
                    row += [c["name"], c["role"], c["ipi"], c["society"], c["per_own"]]
                else:
                    row += ["", "", "", "", ""]
        row += [", ".join(w["issue_codes"]), ", ".join(map(str, w["source_rows"]))]
        rows.append(row)
    return header, rows


def add_step2(wb, iprs):
    header, rows = oneline_table(iprs)
    _sheet(wb, "2 One-line works", header, rows)
    share_labels = ["PER own %", "PER collect %", "MEC own %", "MEC collect %", "SYNC own %", "SYNC collect %"]
    _sheet(wb, "2 Contributors", ["Internal No", "Title", "Excel row", "Set", "Name type", "Role", "Name", "IPI", "Society"] + share_labels,
           [[w["internal_no"], w["title"], c["source_row"], c["set_no"], c["name_type"], c["role"], c["name"], c["ipi"], c["society"]]
            + [c[k] for k in SHARE_KEYS] for w in iprs["works"] for c in w["contributors"]])


# ---- Step 3 -----------------------------------------------------------------
STATE_LABELS = {"linked": "Linked", "review": "Needs review", "unmatched": "Not in master"}


def add_step3(wb, mapping):
    rows = []
    for l in mapping["links"]:
        comp = l.get("comparison") or {}
        diffs = [f"{f['field']}: {f['note'] or (f['iprs'] + ' ≠ ' + f['master'])}" for f in comp.get("fields", [])
                 if f["result"] in ("different", "missing_master")]
        rows.append([l["internal_no"], l["title"], STATE_LABELS[l["state"]], l["method"], l.get("master_row"),
                     l.get("master_title", ""), comp.get("status", l["state"]), "\n".join(diffs) or l["reason"]])
    _sheet(wb, "3 Mapping", ["IPRS Internal No", "IPRS title", "Link", "Linked by", "Master row", "Master title",
                             "Result", "Differences"], rows, status_col=6, widths={"Differences": 80})
    detail = [[l["internal_no"], l["title"], f["field"], f["iprs"], f["master"], f["result"], f["note"]]
              for l in mapping["links"] for f in l.get("comparison", {}).get("fields", [])]
    _sheet(wb, "3 Field comparison", ["IPRS Internal No", "Title", "Field", "IPRS value", "Master value", "Result", "Note"],
           detail, widths={"IPRS value": 50, "Master value": 50, "Note": 60})
    _sheet(wb, "3 Master only", ["Master row", "Title", "IPRS codes in master"],
           [[m["row"], m["title"], ", ".join(m["iprs_codes"])] for m in mapping["master_only"]])


# ---- Step 5 -----------------------------------------------------------------
STATE_NAMES = {"todo": "To do", "in_progress": "In progress", "submitted": "Submitted", "done": "Done", "ignored": "Not needed"}


def add_step5(wb, tunecodes, actions=None):
    actions = actions or {}
    societies = [s for s in ("IPRS", "PRS") if s in tunecodes["summary"]]
    header = ["Title", "Source", "Master row"]
    for s in societies:
        header += [f"{s} tunecode", f"{s} status", f"{s} found by", f"{s} reasons"]
    header += ["Master updates", "Progress"]
    rows = []
    for r in tunecodes["rows"]:
        row = [r["title"], "Added from IPRS" if r["origin"] == "iprs" else "Master", r["master_row"]]
        for s in societies:
            soc = r["societies"][s]
            row += [", ".join(soc["codes"]), soc["status"], soc["method"], "\n".join(soc["reasons"] + soc.get("notes", []))]
        row += ["\n".join(r["master_updates"])]
        targets = [s for s in societies if r["societies"][s]["status"] != "Match"] + (["MASTER"] if r["master_updates"] else [])
        row += ["\n".join(f"{t}: {STATE_NAMES[actions.get(r['key'] + '|' + t, {}).get('state', 'todo')]}" for t in targets)]
        rows.append(row)
    ws = _sheet(wb, "5 Tunecodes", header, rows, widths={f"{s} reasons": 60 for s in societies} | {"Master updates": 60})
    for i, s in enumerate(societies):
        col = 3 + i * 4 + 2
        for n in range(len(rows)):
            fill = STATUS_FILLS.get(rows[n][col - 1])
            if fill:
                ws.cell(n + 2, col).fill = fill
    if tunecodes.get("prs_not_in_report"):
        _sheet(wb, "5 PRS not in report", ["PRS tunecode", "Title", "Writers"],
               [[w["code"], w["title"], ", ".join(w["writers"])] for w in tunecodes["prs_not_in_report"]])


# ---- Step 4: final report in the master's own format -----------------------------
def _put(ws, row, col, value, fill):
    if col is None or value in (None, ""):
        return
    cell = ws.cell(row, col + 1)
    cell.value = value
    cell.fill = fill


def _write_iprs_work(ws, layout, row, w, sn, fill, status):
    """Write one IPRS work as a new row in the master's own (LNV) column layout."""
    fields, slots = layout["fields"], layout["slots"]
    category = [p.strip() for p in w["category"].split("/")]
    values = {"sn": sn, "title": w["title"], "isrc": w["isrc"], "iswc": w["iswc"], "alt_title": w["alt_title"],
              "movie": w["production"], "language": w["language"].title(), "duration": w["duration"],
              "sub_category": category[-1] if category else "", "label": ", ".join(p["name"] for p in w["publishers"])}
    for field, value in values.items():
        _put(ws, row, fields.get(field), value, fill)
    for n, singer in enumerate([s.strip() for s in w["performer"].split(",") if s.strip()][:len(slots.get("singers", []))]):
        _put(ws, row, slots["singers"][n]["name"], singer, fill)
        _put(ws, row, slots["singers"][n]["role"], "S", fill)
    for group in ("composers", "authors", "publishers"):
        people = [c for c in w[group] if not (group == "authors" and c["role"] == "CA")]
        for n, c in enumerate(people[:len(slots.get(group, []))]):
            slot = slots[group][n]
            _put(ws, row, slot["name"], c["name"], fill)
            _put(ws, row, slot["cae"], c["ipi"].lstrip("0"), fill)
            _put(ws, row, slot["role"], c["role"], fill)
            _put(ws, row, slot["perf"], c["per_own"], fill)
            _put(ws, row, slot["mech"], c["mec_own"], fill)
    iprs_blocks = [b for b in layout["blocks"] if b["society"] == "IPRS"]
    if iprs_blocks:
        b = iprs_blocks[0]
        _put(ws, row, b["isrc"], w["isrc"], fill)
        _put(ws, row, b["iswc"], w["iswc"], fill)
        _put(ws, row, b["code"], w["internal_no"], fill)
        _put(ws, row, b["fields"].get("client_status"), status, fill)


def lnv_action_report(master_path, iprs, final, tunecodes, society, status):
    """Our master (LNV format) cut down to the songs that need this action at this society.

    status "Amend"        → songs registered at the society whose registration differs from ours;
    status "Registration" → songs with no tune code at the society.
    The society's block gets Client status = Amendment / Registration and, for amendments, the reasons in
    "Purpose of amendment" (highlighted yellow). A sheet "Why" lists each song and its reasons.
    """
    wb = openpyxl.load_workbook(master_path)
    ws = wb.worksheets[0]
    layout = read_layout([c.value for c in ws[1]])
    blocks = [b for b in layout["blocks"] if b["society"] == society]
    works = {w["internal_no"]: w for w in (iprs or {}).get("works", [])}
    finals = {r["key"]: r for r in final["rows"]}
    label = "Amendment" if status == "Amend" else "Registration"
    picked = [r for r in tunecodes["rows"] if r["societies"].get(society, {}).get("status") == status]

    keep, notes = set(), []
    last = max([s for s in range(1, ws.max_row + 1) if any(ws.cell(s, c).value not in (None, "") for c in (1, 2))] + [1])
    for r in picked:
        res = r["societies"][society]
        if r["master_row"]:
            row = r["master_row"]
        else:  # a song the master does not have yet (added from the IPRS report)
            last += 1
            row = last
            ids = finals[r["key"]]["iprs_works"]
            if ids and ids[0] in works:
                _write_iprs_work(ws, layout, row, works[ids[0]], "", ADDED_FILL, "Added from IPRS report")
        keep.add(row)
        block = next((b for b in blocks if str(ws.cell(row, b["code"] + 1).value or "") in res["codes"]), None) \
            or next((b for b in blocks if ws.cell(row, b["code"] + 1).value in (None, "")), None) or (blocks[0] if blocks else None)
        if block:
            _put(ws, row, block["fields"].get("client_status"), label, FILLED_FILL)
            if status == "Amend":
                _put(ws, row, block["fields"].get("amendment_purpose"), "; ".join(res["reasons"]), FILLED_FILL)
        notes.append([r["title"], ", ".join(res["codes"]) or "—", label, "\n".join(res["reasons"])])

    # Delete the other rows bottom-up, a run of consecutive rows at a time (one delete_rows per row is slow).
    row = ws.max_row
    while row > 1:
        if row in keep:
            row -= 1
            continue
        end = row
        while row > 1 and row not in keep:
            row -= 1
        ws.delete_rows(row + 1, end - row)
    sn_col = layout["fields"].get("sn")
    if sn_col is not None:
        for n in range(2, ws.max_row + 1):
            ws.cell(n, sn_col + 1).value = n - 1

    for name in ("MRM changes", "Why"):
        if name in wb.sheetnames:
            del wb[name]
    why = wb.create_sheet("Why")
    why.append(["Title", f"{society} tune code", "Action", "Reason"])
    for cell in why[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for n in notes:
        why.append(n)
    for col, width in zip("ABCD", (45, 22, 16, 90)):
        why.column_dimensions[col].width = width
    return _bytes(wb)


GROUP_OF = {"Authors": "authors", "Composers": "composers", "Publishers": "publishers"}
SLOT_COLS = ("name", "cae", "joshua", "role", "perf", "mech")


def _master_people(ws, slots, row):
    return [{"name": clean_text(ws.cell(row, s["name"] + 1).value), "ipi": norm_ipi(ws.cell(row, s["cae"] + 1).value)}
            for s in slots if clean_text(ws.cell(row, s["name"] + 1).value)]


def apply_master_edits(ws, layout, edits, works):
    """Write the IPRS value into the master for every field the reviewer chose "IPRS" (or "Add to master") for.
    Returns change-log rows."""
    fields, slots = layout["fields"], layout["slots"]
    changes = []
    chosen = {}
    for e in edits:
        if e["choice"] == "iprs" and e["internal_no"] in works:
            chosen.setdefault(e["master_row"], {})[e["field"]] = works[e["internal_no"]]
    for row, picks in chosen.items():
        for field, w in picks.items():
            detail = f"From IPRS work {w['internal_no']}"
            if field in ("Title", "Duration"):
                value = w["title"] if field == "Title" else w["duration"]
                _put(ws, row, fields.get(field.lower()), value, FILLED_FILL)
            elif field in ("ISWC", "ISRC"):
                key = field.lower()
                value = w[f"{key}_valid"] or w[key]
                _put(ws, row, fields.get(key), value, FILLED_FILL)
                for b in layout["blocks"]:  # the IPRS block that carries this work's tune code
                    if b["society"] == "IPRS" and str(ws.cell(row, b["code"] + 1).value or "").strip() == w["internal_no"]:
                        _put(ws, row, b[key], value, FILLED_FILL)
            else:
                group = GROUP_OF[field]
                people = list(w[group])
                if group == "authors":
                    # A composer/author (CA) sits in the composer slots; list them as author only if the composers
                    # the master will end up with do not already hold them.
                    composers = w["composers"] if "Composers" in picks else _master_people(ws, slots.get("composers", []), row)
                    people = [c for c in people if c["role"] != "CA"
                              or not any(same_party(c["name"], p["name"], c["ipi"], p["ipi"]) for p in composers)]
                group_slots = slots.get(group, [])
                for slot in group_slots:
                    for col in SLOT_COLS:
                        ws.cell(row, slot[col] + 1).value = None
                for n, c in enumerate(people[:len(group_slots)]):
                    slot = group_slots[n]
                    _put(ws, row, slot["name"], c["name"], FILLED_FILL)
                    _put(ws, row, slot["cae"], c["ipi"].lstrip("0"), FILLED_FILL)
                    _put(ws, row, slot["role"], c["role"], FILLED_FILL)
                    _put(ws, row, slot["perf"], c["per_own"], FILLED_FILL)
                    _put(ws, row, slot["mech"], c["mec_own"], FILLED_FILL)
                value = "; ".join(c["name"] for c in people)
                if len(people) > len(group_slots):
                    detail += f" (only {len(group_slots)} of {len(people)} fit the master's slots)"
            title = clean_text(ws.cell(row, fields["title"] + 1).value) if "title" in fields else ""
            changes.append([row, title, f"{field} from IPRS", f"{value} – {detail}"])
    return changes


def _write_fill_row(ws, iprs_blocks, row, f):
    """Write one song's planned IPRS blocks (or its 'Not registered' / 'Not our work' status) into a master row."""
    changes = []
    for b in f["blocks"]:
        block = iprs_blocks[b["slot"] - 1] if b["slot"] <= len(iprs_blocks) else None
        if block is None:
            continue
        _put(ws, row, block["isrc"], b["isrc"], FILLED_FILL)
        _put(ws, row, block["iswc"], b["iswc"], FILLED_FILL)
        _put(ws, row, block["code"], b["internal_no"], FILLED_FILL)
        _put(ws, row, block["fields"].get("client_status"), b["client_status"], FILLED_FILL)
        _put(ws, row, block["fields"].get("general_status"), b["general_status"], FILLED_FILL)
        _put(ws, row, block["fields"].get("amendment_purpose"), b["purpose"], FILLED_FILL)
        changes.append([row, f["title"], f"IPRS block {b['slot']}",
                        f"{b['internal_no']} · Client: {b['client_status']} · General: {b['general_status']}"
                        + (f" · {b['purpose']}" if b["purpose"] else "")])
    if not f["blocks"] and f["state"] in ("not_registered", "not_our_work") and iprs_blocks:
        block = iprs_blocks[0]
        text = "Not our work" if f["state"] == "not_our_work" else "Not registered"
        _put(ws, row, block["fields"].get("client_status"), text, FILLED_FILL)
        _put(ws, row, block["fields"].get("general_status"), text, FILLED_FILL)
        changes.append([row, f["title"], "IPRS block 1", text])
    for code in f["overflow"]:
        changes.append([row, f["title"], "Not filled", f"IPRS work {code}: all {len(iprs_blocks)} IPRS blocks are used"])
    return changes


def final_report(master_path, iprs, final, edits=(), fill=None):
    """Our master, completed: the IPRS columns filled from the IPRS report (yellow), the songs IPRS has that the
    master was missing appended (green), and the values the reviewer chose to take from IPRS (yellow)."""
    wb = openpyxl.load_workbook(master_path)
    ws = wb.worksheets[0]
    layout = read_layout([c.value for c in ws[1]])
    iprs_blocks = [b for b in layout["blocks"] if b["society"] == "IPRS"]
    works = {w["internal_no"]: w for w in iprs["works"]}
    plan = {f["key"]: f for f in (fill or {}).get("rows", [])}
    changes = []

    for r in final["rows"]:
        if r["origin"] == "master" and r["key"] in plan:
            changes += _write_fill_row(ws, iprs_blocks, r["master_row"], plan[r["key"]])

    changes += apply_master_edits(ws, layout, edits, works)

    last = max([s for s in range(1, ws.max_row + 1) if any(ws.cell(s, c).value not in (None, "") for c in (1, 2))] + [1])
    sn = last - 1
    for r in final["added"]:
        w = works[r["iprs_works"][0]]
        last += 1
        sn += 1
        _write_iprs_work(ws, layout, last, w, sn, ADDED_FILL, "")
        changes.append([last, w["title"], "Added song", f"Missing from master; IPRS work {w['internal_no']}"])
        if r["key"] in plan:
            changes += _write_fill_row(ws, iprs_blocks, last, plan[r["key"]])

    if "MRM changes" in wb.sheetnames:
        del wb["MRM changes"]
    log = wb.create_sheet("MRM changes")
    log.append(["Master row", "Title", "Change", "Detail"])
    for cell in log[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for c in changes:
        log.append(c)
    for col, width in zip("ABCD", (12, 45, 22, 90)):
        log.column_dimensions[col].width = width
    return _bytes(wb)


def step_workbook(step, result, actions=None):
    wb = _new_book()
    if step == "step1":
        add_step1(wb, result["iprs"])
    elif step == "step2":
        add_step2(wb, result["iprs"])
    elif step == "step3":
        add_step3(wb, result["mapping"])
    elif step == "step5":
        add_step5(wb, result["tunecodes"], actions)
    elif step == "all":
        if result["iprs"]:
            add_step1(wb, result["iprs"])
            add_step2(wb, result["iprs"])
        if result["mapping"]:
            add_step3(wb, result["mapping"])
        if result["tunecodes"]:
            add_step5(wb, result["tunecodes"], actions)
    return _bytes(wb)


# ---- client-level exports ------------------------------------------------------
def tasks_workbook(client, tasks):
    """Every follow-up for one client, ready to work from or send to a society."""
    wb = _new_book()
    rows = [[t["work"]["mrm_id"], t["work"]["title"], t["society"], t["type_label"], ", ".join(t["codes"]), t["state_label"],
             t["owner"], t["due_date"] or "", "\n".join(t["detail"]), t["note"], t["opened_report"] or "", t["closed_report"] or "",
             t["events"][-1]["text"] if t.get("events") else ""] for t in tasks]
    order = {"Open": 0, "In progress": 1, "Submitted": 2, "Verified in report": 3, "Done": 4, "Dismissed": 5}
    rows.sort(key=lambda r: (order.get(r[5], 9), r[2], r[3], r[1]))
    _sheet(wb, "Tasks", ["MRM ID", "Work", "Where", "Task", "Tune codes", "Status", "Owner", "Due", "Why", "Note",
                         "Found in report", "Closed in report", "Last update"], rows,
           widths={"Why": 70, "Note": 40, "Last update": 50})
    by_society = {}
    for t in tasks:
        if t["state"] not in ("verified", "done", "dismissed"):
            by_society.setdefault((t["society"], t["type_label"]), []).append(t)
    for (society, label), items in sorted(by_society.items()):
        _sheet(wb, f"{society} {label}"[:31], ["MRM ID", "Work", "Tune codes", "Why"],
               [[t["work"]["mrm_id"], t["work"]["title"], ", ".join(t["codes"]), "\n".join(t["detail"])] for t in items],
               widths={"Why": 80})
    return _bytes(wb)


def catalogue_workbook(client, works):
    """Works × society status matrix."""
    wb = _new_book()
    societies = sorted({s for w in works for s in w["registrations"]})
    header = ["MRM ID", "Work", "Source", "ISWC", "IPRS tune codes", "PRS tune codes"] + societies + ["Open tasks", "Client credited"]
    rows = []
    for w in works:
        rows.append([w["mrm_id"], w["title"], "Added from IPRS" if w["origin"] == "iprs" else "Master",
                     ", ".join(w["identifiers"].get("ISWC", [])), ", ".join(w["identifiers"].get("IPRS", [])),
                     ", ".join(w["identifiers"].get("PRS", []))]
                    + [w["registrations"].get(s, {}).get("label", "—") for s in societies]
                    + [w["open_tasks"], "Yes" if w["credited"] else "No"])
    _sheet(wb, "Catalogue", header, rows)
    return _bytes(wb)
